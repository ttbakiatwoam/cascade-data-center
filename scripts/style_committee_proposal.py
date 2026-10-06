#!/usr/bin/env python3
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import sys

path = sys.argv[1]
doc = Document(path)

for sec in doc.sections:
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.32)

styles = doc.styles
styles['Normal'].font.name = 'Liberation Sans'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.08

for sty_name, size, color, before, after in [
    ('Title', 24, '17365D', 0, 12),
    ('Subtitle', 12, '555555', 0, 18),
    ('Heading 1', 17, '17365D', 14, 8),
    ('Heading 2', 13.5, '1F4E79', 12, 6),
    ('Heading 3', 11.5, '2F5597', 10, 4),
]:
    if sty_name in styles:
        s = styles[sty_name]
        s.font.name = 'Liberation Sans'
        s.font.size = Pt(size)
        s.font.bold = sty_name != 'Subtitle'
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

for sec in doc.sections:
    hp = sec.header.paragraphs[0]
    hp.text = 'CASCADE DATA CENTER - DRAFT RECOMMENDATIONS'
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.name = 'Liberation Sans'
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(100,100,100)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('Draft for discussion | October 2026 | Page ')
    r.font.name = 'Liberation Sans'
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100,100,100)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)

for p in doc.paragraphs:
    t = p.text.strip()
    if t.startswith('Appendix A - Proposed Draft Zoning Ordinance'):
        p.paragraph_format.page_break_before = True
    if t.startswith('NON-CODIFIED DRAFTING NOTES'):
        p.paragraph_format.page_break_before = True
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
    if t in ['CITY OF CASCADE, IOWA', 'ORDINANCE NO. ______']:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.font.bold = True

for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, row in enumerate(table.rows):
        if ri == 0:
            trPr = row._tr.get_or_add_trPr()
            tblHeader = OxmlElement('w:tblHeader')
            tblHeader.set(qn('w:val'), 'true')
            trPr.append(tblHeader)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.name = 'Liberation Sans'
                    run.font.size = Pt(9.5)
        if ri == 0:
            for cell in row.cells:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'D9EAF7')
                cell._tc.get_or_add_tcPr().append(shd)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string('17365D')

doc.core_properties.title = 'Cascade Data Center / Major Industrial Facility Zoning Recommendations'
doc.core_properties.subject = 'Draft recommendations for Cascade advisory committee and Planning & Zoning review'
doc.core_properties.author = 'Andrew Saunders'
doc.core_properties.comments = 'Draft for discussion; not adopted City policy or law.'
doc.save(path)
